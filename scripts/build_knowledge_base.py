# -*- coding: utf-8 -*-
"""把 data/raw 中文本层可用的法规文件切成条款块，输出 JSONL 知识库。

输出字段：file, page, clause, title, text, std_no, name, status, impl_date, needs_ocr
"""
import csv
import json
import re
from pathlib import Path

import pdfplumber
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
PROC = ROOT / "data" / "processed"
META_CSV = PROC / "regulations_metadata.csv"
OUT_JSONL = PROC / "knowledge_base.jsonl"

CLAUSE_RE = re.compile(r"^\s*(\d+(?:\.\d+){1,4})\.?\s*(\S.*)?$")
STD_NO_RE = re.compile(
    r"(GB(?:/T|/Z)?\s*[0-9]{4,5}(?:\.\d+)?\s*[-—–]\s*[0-9]{4})|(CQC\s*[0-9]+-[0-9]+-?[0-9]*)",
    re.IGNORECASE,
)

# 乱码/整本扫描件暂时跳过，等 OCR 或字体修复后再入库
DEFER = {
    "GB34660-2017_道路车辆电磁兼容性要求和试验方法.pdf",
    "GBT31484-2015_电动汽车用动力蓄电池循环寿命要求及试验方法.pdf",
    "GB38031-2025_EN_电动汽车用动力蓄电池安全要求.pdf",
}


def load_meta():
    meta = {}
    with META_CSV.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            meta[row["文件名"]] = row
    return meta


def detect_std_no(text):
    for m in STD_NO_RE.finditer(text):
        return re.sub(r"\s+", " ", m.group(0)).strip()
    return None


def enrich(chunk, meta):
    chunk["name"] = meta.get("名称", "")
    chunk["status"] = meta.get("状态", "")
    chunk["impl_date"] = meta.get("实施日期", "")
    chunk.setdefault("std_no", meta.get("标准/规则编号", ""))
    return chunk


def chunk_pdf(path, meta):
    chunks = []
    std_no = meta.get("标准/规则编号") or ""
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                chunks.append({"file": path.name, "page": pno, "clause": None,
                               "title": "", "text": "", "needs_ocr": True})
                continue
            if pno <= 2 and not std_no:
                std_no = detect_std_no(text) or std_no

            cur = None
            preamble = []
            def flush(cur, preamble):
                out = []
                if preamble:
                    out.append({"file": path.name, "page": pno, "clause": None,
                                "title": "", "text": "\n".join(preamble).strip(),
                                "needs_ocr": False})
                    preamble.clear()
                if cur is not None and cur["text"].strip():
                    out.append(cur)
                return out

            for line in text.splitlines():
                m = CLAUSE_RE.match(line)
                if m:
                    chunks.extend(flush(cur, preamble))
                    cur = {"file": path.name, "page": pno,
                           "clause": m.group(1),
                           "title": (m.group(2) or "").strip(" ."),
                           "text": line.strip(), "needs_ocr": False}
                else:
                    if cur is None:
                        preamble.append(line)
                    else:
                        cur["text"] += "\n" + line
            chunks.extend(flush(cur, preamble))
    for c in chunks:
        enrich(c, meta)
        if c.get("std_no"):
            std_no = c["std_no"]
        c["std_no"] = std_no
    return chunks


def chunk_docx(path, meta):
    chunks = []
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    for text in parts:
        chunks.append({"file": path.name, "page": None, "clause": None,
                       "title": "", "text": text, "needs_ocr": False,
                       "std_no": meta.get("标准/规则编号", "")})
    for c in chunks:
        enrich(c, meta)
    return chunks


def main():
    meta = load_meta()
    all_chunks = []
    for path in sorted(RAW.iterdir()):
        if path.name in DEFER:
            continue
        m = meta.get(path.name, {})
        if path.suffix.lower() == ".pdf":
            all_chunks.extend(chunk_pdf(path, m))
        elif path.suffix.lower() == ".docx":
            all_chunks.extend(chunk_docx(path, m))
    OUT_JSONL.write_text(
        "\n".join(json.dumps(c, ensure_ascii=False) for c in all_chunks) + "\n",
        encoding="utf-8",
    )
    print(f"chunks={len(all_chunks)} -> {OUT_JSONL}")


if __name__ == "__main__":
    main()
