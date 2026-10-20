# -*- coding: utf-8 -*-
import json
import re
from pathlib import Path

import pdfplumber
from docx import Document

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
SAMPLES = ROOT / "data" / "raw"
OUT_JSON = ROOT / "outputs" / "parse_report.json"

CLAUSE_RE = re.compile(r"^\s*(\d+(?:\.\d+){1,4})\s*(\S.*)?$")
CHAPTER_RE = re.compile(r"^\s*(第[一二三四五六七八九十百千0-9]+[章条部分])")
ANNEX_RE = re.compile(r"^\s*(附录\s*[A-Z]?|附\s*录\s*[A-Z])", re.IGNORECASE)
TABLE_RE = re.compile(r"^\s*(表\s*[A-Z]?\s*\d+(?:\.\d+)?|表格\s*\d+)", re.IGNORECASE)
STD_NO_RE = re.compile(r"(GB(?:/T|/Z)?\s*[0-9]{4,5}(?:\.\d+)?\s*[-—–]\s*[0-9]{4})|(CQC\s*[0-9]+-[0-9]+-?[0-9]*)", re.IGNORECASE)
DATE_RE = re.compile(r"(20[12][0-9])\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")

def weird_ratio(text):
    if not text:
        return 0.0
    n = 0
    for ch in text:
        o = ord(ch)
        if 0xE000 <= o <= 0xF8FF or 0xFFF0 <= o <= 0xFFFF or o == 0xFFFD or 0x0080 <= o <= 0x009F:
            n += 1
    return n / len(text)

def analyze_pdf(path):
    rec = {"file": path.name, "type": "pdf", "pages": 0, "text_chars": 0,
           "no_text_pages": [], "low_text_pages": [], "suspect_encoding_pages": [],
           "clause_hits": 0, "chapter_hits": 0, "annex_hits": 0, "table_hits": 0,
           "std_numbers": [], "dates": [], "error": None, "sample_clauses": []}
    try:
        with pdfplumber.open(path) as pdf:
            rec["pages"] = len(pdf.pages)
            for i, page in enumerate(pdf.pages, 1):
                try:
                    text = page.extract_text() or ""
                except Exception as e:
                    text = ""
                    rec["error"] = (rec["error"] or "") + f"p{i}:{e}; "
                clean = text.replace("\n", " ").strip()
                rec["text_chars"] += len(clean)
                if not clean:
                    rec["no_text_pages"].append(i)
                elif len(clean) < 20:
                    rec["low_text_pages"].append(i)
                if clean and weird_ratio(clean) > 0.05:
                    rec["suspect_encoding_pages"].append(i)
                for line in text.splitlines():
                    if CLAUSE_RE.match(line):
                        rec["clause_hits"] += 1
                        if len(rec["sample_clauses"]) < 12:
                            rec["sample_clauses"].append(line.strip()[:80])
                    if CHAPTER_RE.match(line):
                        rec["chapter_hits"] += 1
                    if ANNEX_RE.match(line):
                        rec["annex_hits"] += 1
                    if TABLE_RE.match(line):
                        rec["table_hits"] += 1
                for m in STD_NO_RE.finditer(text):
                    v = re.sub(r"\s+", " ", m.group(0)).strip()
                    if v and v not in rec["std_numbers"]:
                        rec["std_numbers"].append(v)
                for m in DATE_RE.finditer(text):
                    v = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
                    if v not in rec["dates"]:
                        rec["dates"].append(v)
    except Exception as e:
        rec["error"] = str(e)
    return rec

def analyze_docx(path):
    rec = {"file": path.name, "type": "docx", "pages": None, "text_chars": 0,
           "paragraphs": 0, "tables": 0, "table_rows": 0, "clause_hits": 0,
           "std_numbers": [], "dates": [], "error": None, "sample_clauses": []}
    try:
        doc = Document(path)
        text_parts = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                text_parts.append(t)
                rec["paragraphs"] += 1
        rec["tables"] = len(doc.tables)
        for table in doc.tables:
            for row in table.rows:
                rec["table_rows"] += 1
                row_text = " ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    text_parts.append(row_text)
        full = "\n".join(text_parts)
        rec["text_chars"] = len(full.replace("\n", ""))
        for line in text_parts:
            if CLAUSE_RE.match(line):
                rec["clause_hits"] += 1
                if len(rec["sample_clauses"]) < 12:
                    rec["sample_clauses"].append(line[:80])
        for m in STD_NO_RE.finditer(full):
            v = re.sub(r"\s+", " ", m.group(0)).strip()
            if v and v not in rec["std_numbers"]:
                rec["std_numbers"].append(v)
        for m in DATE_RE.finditer(full):
            v = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
            if v not in rec["dates"]:
                rec["dates"].append(v)
    except Exception as e:
        rec["error"] = str(e)
    return rec

records = []
for path in sorted(SAMPLES.iterdir()):
    if path.suffix.lower() == ".pdf":
        records.append(analyze_pdf(path))
    elif path.suffix.lower() == ".docx":
        records.append(analyze_docx(path))

OUT_JSON.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"analyzed {len(records)} files -> {OUT_JSON}")
